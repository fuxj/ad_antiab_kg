#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
08d_export_xlsx.py  --  chapter-6 final table export (POST-PROCESSING ONLY)

This script does NOT touch the fusion logic and does NOT recompute any score.
It reads the paper tables already produced by 08c_tables.py and:

  1. tags the directional false positive NIROGACESTAT.  It entered the Task-A
     list only because its target (PSEN1 / PSEN2) is on the panel, but PSEN1
     sits in the CONTROL stratum (no MR support) and gamma-secretase
     inhibition worsened cognition in phase-3 trials.  It is therefore
     demoted from "candidate" to a worked example of how target matching
     without direction can mislead.  The row is KEPT (it is worth discussing)
     but moved to the tail of the table and flagged.
  2. guarantees the acetaminophen footnote is present in the Task-B table.
  3. writes a formatted Excel workbook (one sheet per table) so the
     supplementary tables can be edited directly in Excel.
  4. emits core_risk_drugs.csv -- the systemic first-line / high-tier
     co-medication shortlist that the manuscript actually discusses.

Reads  (under <proj>/results/paper_tables/):
    Table_S1_repositioning.csv
    Table_S2_comedication_liability.csv
    Table_S3_combination_partners.csv
    dedup_and_class_audit.csv

Writes (under <proj>/results/paper_tables/):
    Table_S1_repositioning.csv   (in place, plus the new tag columns)
    paper_tables.xlsx
    core_risk_drugs.csv
    xlsx_build_summary.json
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd

VERSION = "1.0.0"

# name fragment used to locate the directional false positive
FP_MATCH = "NIROGACESTAT"
FP_NOTE = (
    "directional false positive: gamma-secretase inhibitor; PSEN1 sits in the "
    "control stratum (no MR support) and the class worsened cognition in phase 3"
)
ACETAMINOPHEN_NOTE = (
    "known false positive: only weak antiplatelet activity; entered via PTGS1 target"
)

# preferred reading order of the candidate roles in the Task-A table
ROLE_ORDER = [
    "genetic-anchored",
    "genetic-counter",
    "target-prior",
    "structural-only",
    "false-positive-direction",
]
ROLE_RANK = {r: i for i, r in enumerate(ROLE_ORDER)}

TAG_COLS = ["candidate_role", "directional_false_positive", "fp_note", "rank_final"]

# sheet name -> (column to colour, low colour, high colour)
COLOUR_PLAN = {
    "S1 Repurposing": ("f_repurp", "FFFFFF", "4472C4"),
    "S2 CoMed Liability": ("liability_adj", "FFFFFF", "F8696B"),
    "S3 Partners": ("f_partner", "FFFFFF", "63BE7B"),
}


def log(msg: str) -> None:
    print(f"[{datetime.now():%H:%M:%S}] {msg}", flush=True)


def num(s) -> pd.Series:
    return pd.to_numeric(s, errors="coerce")


def read_csv(p: Path):
    p = Path(p)
    if not p.exists():
        log(f"[warn] missing: {p.name}")
        return None
    try:
        return pd.read_csv(p, encoding="utf-8-sig")
    except Exception as e:  # noqa: BLE001
        log(f"[warn] cannot read {p.name}: {e}")
        return None


def _haystack(df: pd.DataFrame) -> pd.Series:
    """upper-case 'drug + active_moiety' string used for name matching."""
    a = df["drug"].astype(str) if "drug" in df.columns else pd.Series("", index=df.index)
    b = (df["active_moiety"].astype(str) if "active_moiety" in df.columns
         else pd.Series("", index=df.index))
    return (a + " " + b).str.upper()


# --------------------------------------------------------------------------- #
def tag_s1(s1: pd.DataFrame) -> pd.DataFrame:
    """Attach candidate_role / false-positive tags, then re-order the table."""
    s1 = s1.copy()
    s1 = s1.drop(columns=[c for c in TAG_COLS if c in s1.columns])

    aligned = num(s1["ad_eff_aligned"]) if "ad_eff_aligned" in s1.columns else pd.Series(0.0, index=s1.index)
    counter = num(s1["ad_eff_counter"]) if "ad_eff_counter" in s1.columns else pd.Series(0.0, index=s1.index)
    prior = num(s1["n_prior_targets"]) if "n_prior_targets" in s1.columns else pd.Series(0.0, index=s1.index)
    gp = s1["gene_prior"] if "gene_prior" in s1.columns else pd.Series(np.nan, index=s1.index)

    def role_of(i) -> str:
        if float(aligned.get(i, 0.0) or 0.0) > 0:
            return "genetic-anchored"
        if float(counter.get(i, 0.0) or 0.0) > 0:
            return "genetic-counter"
        if float(prior.get(i, 0.0) or 0.0) > 0 or pd.notna(gp.get(i)):
            return "target-prior"
        return "structural-only"

    s1["candidate_role"] = [role_of(i) for i in s1.index]
    s1["directional_false_positive"] = 0
    s1["fp_note"] = ""

    hit = _haystack(s1).str.contains(FP_MATCH)
    if bool(hit.any()):
        s1.loc[hit, "candidate_role"] = "false-positive-direction"
        s1.loc[hit, "directional_false_positive"] = 1
        s1.loc[hit, "fp_note"] = FP_NOTE
        log(f"[tag] directional false positive flagged: "
            f"{', '.join(sorted(set(s1.loc[hit, 'drug'].astype(str))))}")
    else:
        log(f"[warn] '{FP_MATCH}' not found in Table_S1 -> nothing demoted")

    order = s1["candidate_role"].map(ROLE_RANK).fillna(99).astype(int)
    if "rank_paper" in s1.columns:
        orig = num(s1["rank_paper"]).fillna(1e9)
    else:
        orig = pd.Series(np.arange(1, len(s1) + 1), index=s1.index)
    s1 = (s1.assign(_ro=order, _orig=orig)
            .sort_values(["_ro", "_orig"])
            .drop(columns=["_ro", "_orig"])
            .reset_index(drop=True))
    s1["rank_final"] = np.arange(1, len(s1) + 1)
    return s1


def ensure_acetaminophen_note(s2: pd.DataFrame) -> pd.DataFrame:
    s2 = s2.copy()
    if "paper_note" not in s2.columns:
        s2["paper_note"] = ""
    s2["paper_note"] = s2["paper_note"].fillna("").astype(str)
    hit = _haystack(s2).str.contains("ACETAMINOPHEN")
    blank = s2["paper_note"].str.strip() == ""
    m = hit & blank
    if bool(m.any()):
        s2.loc[m, "paper_note"] = ACETAMINOPHEN_NOTE
        log(f"[tag] acetaminophen footnote restored on {int(m.sum())} row(s)")
    return s2


def core_risk(s2: pd.DataFrame):
    if s2 is None:
        return None
    cls = s2["clinical_class"].astype(str) if "clinical_class" in s2.columns else pd.Series("", index=s2.index)
    tier = s2["liability_tier"].astype(str) if "liability_tier" in s2.columns else pd.Series("", index=s2.index)
    m = (cls == "systemic_first_line") & (tier == "high")
    w = s2[m].copy()
    if "liability_adj" in w.columns:
        w = w.sort_values("liability_adj", ascending=False).reset_index(drop=True)
    return w


# --------------------------------------------------------------------------- #
def write_xlsx(path: Path, sheets) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets:
            if df is None:
                df = pd.DataFrame()
            df.to_excel(xw, sheet_name=name[:31], index=False)

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        # column widths (sampled, capped)
        for ci, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1):
            m = 0
            for cell in col:
                if cell.value is not None:
                    m = max(m, len(str(cell.value)))
            ws.column_dimensions[get_column_letter(ci)].width = min(46, max(9, m + 2))

        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions

        plan = COLOUR_PLAN.get(ws.title)
        if plan and ws.max_row > 1:
            key, low, high = plan
            header = [c.value for c in ws[1]]
            if key in header:
                ci = header.index(key) + 1
                col = get_column_letter(ci)
                rng = f"{col}2:{col}{ws.max_row}"
                ws.conditional_formatting.add(
                    rng,
                    ColorScaleRule(start_type="num", start_value=0.0, start_color=low,
                                   end_type="num", end_value=1.0, end_color=high))
    wb.save(path)


# --------------------------------------------------------------------------- #
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="chapter-6 xlsx export + FP tagging")
    ap.add_argument("--proj", default=".")
    args = ap.parse_args(argv)

    proj = Path(args.proj).resolve()
    OUT = proj / "results" / "paper_tables"
    if not OUT.exists():
        log(f"[error] {OUT} not found -> run 08c_tables.py first")
        return 2

    log(f"project root: {proj}")
    log(f"=== STAGE table export (v{VERSION}) ===")

    summary = {"version": VERSION, "output_dir": str(OUT), "sheets": {}}

    s1 = read_csv(OUT / "Table_S1_repositioning.csv")
    s2 = read_csv(OUT / "Table_S2_comedication_liability.csv")
    s3 = read_csv(OUT / "Table_S3_combination_partners.csv")
    audit = read_csv(OUT / "dedup_and_class_audit.csv")

    # -- S1: role tagging + false-positive demotion ------------------------- #
    if s1 is not None:
        n0 = len(s1)
        s1 = tag_s1(s1)
        s1.to_csv(OUT / "Table_S1_repositioning.csv", index=False)
        role_counts = s1["candidate_role"].value_counts().to_dict()
        nfp = int((s1["directional_false_positive"] == 1).sum())
        log(f"[write] Table_S1_repositioning.csv | {n0} rows | roles={role_counts} | FP={nfp}")
        summary["sheets"]["S1_repositioning"] = {
            "rows": int(len(s1)), "roles": {k: int(v) for k, v in role_counts.items()},
            "false_positive_tagged": nfp}
    else:
        log("[warn] Table_S1 missing -> skipped")

    # -- S2: acetaminophen footnote ----------------------------------------- #
    core = None
    if s2 is not None:
        s2 = ensure_acetaminophen_note(s2)
        s2.to_csv(OUT / "Table_S2_comedication_liability.csv", index=False)
        core = core_risk(s2)
        if core is not None:
            core.to_csv(OUT / "core_risk_drugs.csv", index=False)
            log(f"[write] core_risk_drugs.csv | {len(core)} systemic first-line high-tier drug(s)")
        summary["sheets"]["S2_comedication"] = {
            "rows": int(len(s2)),
            "core_risk_rows": int(0 if core is None else len(core))}
    else:
        log("[warn] Table_S2 missing -> skipped")

    # -- workbook ----------------------------------------------------------- #
    xlsx = OUT / "paper_tables.xlsx"
    sheets = [
        ("S1 Repurposing", s1),
        ("S2 CoMed Liability", s2),
        ("S3 Partners", s3),
        ("Audit", audit),
    ]
    if core is not None and len(core) > 0:
        sheets.append(("Core Risk Drugs", core))
    write_xlsx(xlsx, sheets)
    log(f"[write] {xlsx.name} | {len(sheets)} sheet(s): "
        f"{', '.join(n for n, _ in sheets)}")

    if s3 is not None:
        summary["sheets"]["S3_partners"] = {"rows": int(len(s3))}
    if core is not None:
        log("--- core high-risk co-medication shortlist (top 15) ---")
        cols = [c for c in ["rank_paper", "drug", "liability_adj", "clinical_group",
                            "bleed_class", "bleed_genes", "paper_note"] if c in core.columns]
        for _, r in core.head(15).iterrows():
            note = str(r.get("paper_note", ""))[:40]
            log(f"   {str(r.get('drug')):26s} adj={float(r.get('liability_adj', float('nan'))):.3f} "
                f"grp={str(r.get('clinical_group')):20s} {note}")

    summary["generated_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    (OUT / "xlsx_build_summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    log("done.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
